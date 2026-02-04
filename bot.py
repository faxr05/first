"""
Telegram Bot - Presentation & Report Generator
SQLite DATABASE - Ishonchli va sodda
"""

import os
import json
import asyncio
import aiosqlite
from datetime import date
from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, FSInputFile
import google.generativeai as genai
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from docx import Document

# ============ KONFIGURATSIYA ============
BOT_TOKEN = os.getenv('BOT_TOKEN', 'YOUR_BOT_TOKEN')
GEMINI_API_KEY = os.getenv('GEMINI_API_KEY', 'YOUR_GEMINI_KEY')
REQUIRED_CHANNEL = ('@bkzsdfgahd')
DB_PATH = 'bot_data.db'

# Gemini AI
genai.configure(api_key=GEMINI_API_KEY)
model = genai.GenerativeModel('gemini-2.5-flash')

# Bot
bot = Bot(token=BOT_TOKEN)
storage = MemoryStorage()
dp = Dispatcher(storage=storage)

# ============ DATABASE ============
async def init_db():
    """Database yaratish"""
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("""
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                username TEXT,
                first_name TEXT,
                language TEXT DEFAULT 'uz',
                daily_limit INTEGER DEFAULT 2,
                used_today INTEGER DEFAULT 0,
                last_reset TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        await db.commit()

async def get_user(user_id: int):
    """User ma'lumotlarini olish"""
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)) as cursor:
            row = await cursor.fetchone()
            return dict(row) if row else None

async def add_user(user_id: int, username: str = None, first_name: str = None):
    """Yangi user qo'shish"""
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("""
            INSERT OR IGNORE INTO users (user_id, username, first_name, last_reset)
            VALUES (?, ?, ?, ?)
        """, (user_id, username, first_name, str(date.today())))
        await db.commit()

async def update_language(user_id: int, language: str):
    """Til yangilash"""
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("UPDATE users SET language = ? WHERE user_id = ?", (language, user_id))
        await db.commit()

async def check_and_reset_limit(user_id: int) -> tuple:
    """Limitni tekshirish va reset"""
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute("SELECT daily_limit, used_today, last_reset FROM users WHERE user_id = ?", (user_id,)) as cursor:
            row = await cursor.fetchone()
            
            if not row:
                return (2, 2)
            
            today = str(date.today())
            if row['last_reset'] != today:
                await db.execute("UPDATE users SET used_today = 0, last_reset = ? WHERE user_id = ?", (today, user_id))
                await db.commit()
                return (row['daily_limit'], row['daily_limit'])
            
            remaining = row['daily_limit'] - row['used_today']
            return (remaining, row['daily_limit'])

async def use_generation(user_id: int):
    """Generatsiyani ishlatish"""
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("UPDATE users SET used_today = used_today + 1 WHERE user_id = ?", (user_id,))
        await db.commit()

# ============ HOLATLAR ============
class BotStates(StatesGroup):
    lang_select = State()
    check_subscription = State()
    select_type = State()
    enter_topic = State()
    enter_pages = State()
    select_design = State()
    confirm = State()

# ============ TARJIMALAR ============
TEXTS = {
    'uz': {
        'welcome': '👋 Assalomu alaykum!\n\n📊 Taqdimot va 📝 Referat/Mustaqil ish tayyorlash botiga xush kelibsiz!\n\n💎 Kunlik limit: {remaining}/{total}\n\nTilni tanlang:',
        'subscription_required': '📢 Botdan foydalanish uchun kanalga obuna bo\'ling:\n\n⛔️{channel}⛔️\n\n✅ Obuna bo\'lgach "Tasdiqlash" tugmasini bosing',
        'check_btn': '✅ Obuna tekshirish',
        'not_subscribed': '❌ Siz hali obuna bo\'lmadingiz!\n\nIltimos, avval kanalga obuna bo\'ling: ⛔️{channel}⛔️',
        'select_type': '📑 Qaysi turdagi hujjat kerak?\n\n💎 Bugungi limit: {remaining}/{total}',
        'presentation': '📊 Taqdimot (PPTX)',
        'report': '📝 Referat',
        'coursework': '📚 Mustaqil ish',
        'enter_topic': '✍️ Mavzuni kiriting:',
        'enter_pages': '📄 Nechta sahifa kerak? (5-20 oralig\'ida)',
        'invalid_pages': '❌ Noto\'g\'ri son! 5 dan 20 gacha son kiriting.',
        'select_design': '🎨 Dizayn shablonini tanlang:',
        'confirm_data': '📋 <b>Kiritilgan ma\'lumotlar:</b>\n\n🎯 Tur: {doc_type}\n📖 Mavzu: {topic}\n📄 Sahifalar: {pages}\n{design}\n✅ Davom etamizmi?',
        'confirm_yes': '✅ Ha, davom etish',
        'confirm_no': '❌ Yo\'q, qaytadan',
        'generating': '⏳ Tayyorlanmoqda... Iltimos kuting...\n\n📊 Bu 30-60 soniya vaqt olishi mumkin.',
        'success': '✅ Tayyor! Marhamat:\n\n💎 Qolgan limit: {remaining}/{total}',
        'error': '❌ Xatolik yuz berdi. Qaytadan urinib ko\'ring.',
        'back_to_start': '🔙 Boshiga qaytish',
        'limit_reached': '⛔️ Kunlik limitingiz tugadi!\n\n💎 Limit: {remaining}/{total}\n🔄 Ertaga yangi limit beriladi'
    },
    'ru': {
        'welcome': '👋 Здравствуйте!\n\n📊 Бот для создания презентаций и 📝 рефератов!\n\n💎 Лимит: {remaining}/{total}\n\nВыберите язык:',
        'subscription_required': '📢 Подпишитесь на канал:\n\n⛔️{channel}⛔️\n\n✅ Нажмите "Проверить"',
        'check_btn': '✅ Проверить',
        'not_subscribed': '❌ Вы не подписались!\n\nПодпишитесь: ⛔️{channel}⛔️',
        'select_type': '📑 Тип документа?\n\n💎 Лимит: {remaining}/{total}',
        'presentation': '📊 Презентация',
        'report': '📝 Реферат',
        'coursework': '📚 Курсовая',
        'enter_topic': '✍️ Введите тему:',
        'enter_pages': '📄 Страниц? (5-20)',
        'invalid_pages': '❌ Неверно! 5-20.',
        'select_design': '🎨 Дизайн:',
        'confirm_data': '📋 <b>Данные:</b>\n\n🎯 Тип: {doc_type}\n📖 Тема: {topic}\n📄 Страниц: {pages}\n{design}\n✅ Продолжить?',
        'confirm_yes': '✅ Да',
        'confirm_no': '❌ Нет',
        'generating': '⏳ Генерация... Подождите...\n\n📊 30-60 секунд.',
        'success': '✅ Готово!\n\n💎 Осталось: {remaining}/{total}',
        'error': '❌ Ошибка. Попробуйте снова.',
        'back_to_start': '🔙 Назад',
        'limit_reached': '⛔️ Лимит исчерпан!\n\n💎 Лимит: {remaining}/{total}\n🔄 Завтра новый'
    },
    'en': {
        'welcome': '👋 Hello!\n\n📊 Presentation & Report Bot!\n\n💎 Limit: {remaining}/{total}\n\nSelect language:',
        'subscription_required': '📢 Subscribe:\n\n⛔️{channel}⛔️\n\n✅ Click "Check"',
        'check_btn': '✅ Check',
        'not_subscribed': '❌ Not subscribed!\n\nSubscribe: ⛔️{channel}⛔️',
        'select_type': '📑 Document type?\n\n💎 Limit: {remaining}/{total}',
        'presentation': '📊 Presentation',
        'report': '📝 Report',
        'coursework': '📚 Coursework',
        'enter_topic': '✍️ Topic:',
        'enter_pages': '📄 Pages? (5-20)',
        'invalid_pages': '❌ Invalid! 5-20.',
        'select_design': '🎨 Design:',
        'confirm_data': '📋 <b>Data:</b>\n\n🎯 Type: {doc_type}\n📖 Topic: {topic}\n📄 Pages: {pages}\n{design}\n✅ Continue?',
        'confirm_yes': '✅ Yes',
        'confirm_no': '❌ No',
        'generating': '⏳ Generating...\n\n📊 30-60 sec.',
        'success': '✅ Done!\n\n💎 Left: {remaining}/{total}',
        'error': '❌ Error. Try again.',
        'back_to_start': '🔙 Back',
        'limit_reached': '⛔️ Limit reached!\n\n💎 Limit: {remaining}/{total}\n🔄 Tomorrow: new'
    }
}

DESIGNS = {
    '1': {'name': 'Klassik Ko\'k', 'bg': (31, 78, 121), 'title': (255, 255, 255), 'text': (0, 0, 0)},
    '2': {'name': 'Professional', 'bg': (68, 114, 196), 'title': (255, 255, 255), 'text': (0, 0, 0)},
    '3': {'name': 'Zamonaviy', 'bg': (91, 155, 213), 'title': (255, 255, 255), 'text': (0, 0, 0)},
    '4': {'name': 'Qizil', 'bg': (192, 0, 0), 'title': (255, 255, 255), 'text': (0, 0, 0)},
    '5': {'name': 'Yashil', 'bg': (0, 176, 80), 'title': (255, 255, 255), 'text': (0, 0, 0)}
}

# ============ FUNKSIYALAR ============
async def check_subscription(user_id: int) -> bool:
    try:
        member = await bot.get_chat_member(chat_id=REQUIRED_CHANNEL, user_id=user_id)
        return member.status in ['member', 'administrator', 'creator']
    except:
        return False

def get_text(lang: str, key: str) -> str:
    return TEXTS.get(lang, TEXTS['uz']).get(key, '')

async def generate_content_with_gemini(topic: str, pages: int, doc_type: str, lang: str) -> dict:
    lang_map = {'uz': 'uzbek', 'ru': 'russian', 'en': 'english'}
    lang_full = lang_map.get(lang, 'uzbek')
    
    if doc_type == 'presentation':
        prompt = f"""Create presentation in {lang_full} about "{topic}". Generate EXACTLY {pages} slides.
Return ONLY valid JSON:
{{"title": "Main title", "slides": [{{"title": "Slide 1", "content": ["Point 1", "Point 2", "Point 3"]}}]}}
Requirements: Each slide 5-7 points, {lang_full} language, ONLY JSON"""
    else:
        prompt = f"""Create {'report' if doc_type == 'report' else 'coursework'} in {lang_full} about "{topic}". For {pages} pages.
Return ONLY valid JSON:
{{"title": "Title", "introduction": "Intro (4-5 paragraphs)", "sections": [{{"title": "Section 1", "content": "Content (4-5 paragraphs)"}}], "conclusion": "Conclusion (2-3 paragraphs)"}}
Requirements: Enough sections for {pages} pages, {lang_full}, ONLY JSON"""

    try:
        response = model.generate_content(prompt)
        result_text = response.text.strip()
        
        if result_text.startswith('```'):
            result_text = result_text.split('```')[1]
            if result_text.startswith('json'):
                result_text = result_text[4:]
            result_text = result_text.strip()
        
        return json.loads(result_text)
    except:
        raise Exception("AI javobini tahlil qilishda xatolik")

def create_presentation(data: dict, design_id: str, output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    design = DESIGNS[design_id]
    bg_color = design['bg']
    title_color = design['title']
    text_color = design['text']
    
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)
    
    title_box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = data['title']
    title_frame.paragraphs[0].font.size = Pt(44)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(*title_color)
    title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    for slide_data in data['slides']:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(255, 255, 255)
        
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
        title_frame = title_box.text_frame
        title_frame.text = slide_data['title']
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(*bg_color)
        
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5))
        text_frame = content_box.text_frame
        text_frame.word_wrap = True
        
        for point in slide_data['content']:
            p = text_frame.add_paragraph()
            p.text = f"• {point}"
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(*text_color)
            p.space_before = Pt(12)
    
    prs.save(output_path)

def create_document(data: dict, output_path: str):
    doc = Document()
    
    title = doc.add_heading(data['title'], 0)
    title.alignment = 1
    
    doc.add_heading('Kirish', 1)
    doc.add_paragraph(data['introduction'])
    
    for section in data['sections']:
        doc.add_heading(section['title'], 1)
        doc.add_paragraph(section['content'])
    
    doc.add_heading('Xulosa', 1)
    doc.add_paragraph(data['conclusion'])
    
    doc.save(output_path)

# ============ HANDLERLAR ============
@dp.message(Command("start"))
async def cmd_start(message: types.Message, state: FSMContext):
    await state.clear()
    user_id = message.from_user.id
    username = message.from_user.username
    first_name = message.from_user.first_name
    
    user = await get_user(user_id)
    if not user:
        await add_user(user_id, username, first_name)
    
    remaining, total = await check_and_reset_limit(user_id)
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🇺🇿 O'zbek", callback_data="lang_uz")],
        [InlineKeyboardButton(text="🇷🇺 Русский", callback_data="lang_ru")],
        [InlineKeyboardButton(text="🇬🇧 English", callback_data="lang_en")]
    ])
    
    text = f"👋 Assalomu alaykum! | Здравствуйте! | Hello!\n\n📊 Taqdimot va 📝 Referat boti\n\n💎 Limit: {remaining}/{total}\n\nTilni tanlang:"
    
    await message.answer(text, reply_markup=keyboard)
    await state.set_state(BotStates.lang_select)

@dp.callback_query(F.data.startswith("lang_"))
async def process_language(callback: types.CallbackQuery, state: FSMContext):
    lang = callback.data.split("_")[1]
    await state.update_data(lang=lang)
    await update_language(callback.from_user.id, lang)
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=get_text(lang, 'check_btn'), callback_data="check_sub")]
    ])
    
    text = get_text(lang, 'subscription_required').format(channel=REQUIRED_CHANNEL)
    await callback.message.edit_text(text, reply_markup=keyboard)
    await state.set_state(BotStates.check_subscription)
    await callback.answer()

@dp.callback_query(F.data == "check_sub")
async def check_sub(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    lang = data.get('lang', 'uz')
    user_id = callback.from_user.id
    
    is_subscribed = await check_subscription(user_id)
    
    if is_subscribed:
        remaining, total = await check_and_reset_limit(user_id)
        
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=get_text(lang, 'presentation'), callback_data="type_presentation")],
            [InlineKeyboardButton(text=get_text(lang, 'report'), callback_data="type_report")],
            [InlineKeyboardButton(text=get_text(lang, 'coursework'), callback_data="type_coursework")]
        ])
        
        text = get_text(lang, 'select_type').format(remaining=remaining, total=total)
        await callback.message.edit_text(text, reply_markup=keyboard)
        await state.set_state(BotStates.select_type)
    else:
        await callback.answer(get_text(lang, 'not_subscribed').format(channel=REQUIRED_CHANNEL), show_alert=True)
    await callback.answer()

@dp.callback_query(F.data.startswith("type_"))
async def select_type(callback: types.CallbackQuery, state: FSMContext):
    user_id = callback.from_user.id
    data = await state.get_data()
    lang = data.get('lang', 'uz')
    
    remaining, total = await check_and_reset_limit(user_id)
    
    if remaining <= 0:
        text = get_text(lang, 'limit_reached').format(remaining=remaining, total=total)
        await callback.message.edit_text(text)
        await callback.answer()
        return
    
    doc_type = callback.data.split("_")[1]
    await state.update_data(doc_type=doc_type)
    
    await callback.message.edit_text(get_text(lang, 'enter_topic'))
    await state.set_state(BotStates.enter_topic)
    await callback.answer()

@dp.message(BotStates.enter_topic)
async def enter_topic(message: types.Message, state: FSMContext):
    await state.update_data(topic=message.text)
    data = await state.get_data()
    lang = data.get('lang', 'uz')
    
    await message.answer(get_text(lang, 'enter_pages'))
    await state.set_state(BotStates.enter_pages)

@dp.message(BotStates.enter_pages)
async def enter_pages(message: types.Message, state: FSMContext):
    data = await state.get_data()
    lang = data.get('lang', 'uz')
    
    try:
        pages = int(message.text)
        if pages < 5 or pages > 20:
            await message.answer(get_text(lang, 'invalid_pages'))
            return
        
        await state.update_data(pages=pages)
        doc_type = data.get('doc_type')
        
        if doc_type == 'presentation':
            designs_text = "🎨 <b>Dizaynlar:</b>\n\n"
            keyboard_buttons = []
            
            for design_id, design_info in DESIGNS.items():
                #designs_text += f"{design_id}. {design_info['name']}\n"
                keyboard_buttons.append([InlineKeyboardButton(text=f"{design_id}. {design_info['name']}", callback_data=f"design_{design_id}")])
            
            keyboard = InlineKeyboardMarkup(inline_keyboard=keyboard_buttons)
            await message.answer(designs_text, reply_markup=keyboard, parse_mode='HTML')
            await state.set_state(BotStates.select_design)
        else:
            await show_confirmation(message, state)
    except ValueError:
        await message.answer(get_text(lang, 'invalid_pages'))

@dp.callback_query(F.data.startswith("design_"))
async def select_design(callback: types.CallbackQuery, state: FSMContext):
    design_id = callback.data.split("_")[1]
    await state.update_data(design=design_id)
    await show_confirmation(callback.message, state)
    await callback.answer()

async def show_confirmation(message: types.Message, state: FSMContext):
    data = await state.get_data()
    lang = data.get('lang', 'uz')
    doc_type = data.get('doc_type')
    topic = data.get('topic')
    pages = data.get('pages')
    design = data.get('design')
    
    type_names = {
        'uz': {'presentation': 'Taqdimot', 'report': 'Referat', 'coursework': 'Mustaqil ish'},
        'ru': {'presentation': 'Презентация', 'report': 'Реферат', 'coursework': 'Курсовая'},
        'en': {'presentation': 'Presentation', 'report': 'Report', 'coursework': 'Coursework'}
    }
    
    doc_type_name = type_names[lang][doc_type]
    design_text = f"🎨 Dizayn: {DESIGNS[design]['name']}\n" if design else ""
    
    text = get_text(lang, 'confirm_data').format(doc_type=doc_type_name, topic=topic, pages=pages, design=design_text)
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=get_text(lang, 'confirm_yes'), callback_data="confirm_yes")],
        [InlineKeyboardButton(text=get_text(lang, 'confirm_no'), callback_data="confirm_no")]
    ])
    
    await message.answer(text, reply_markup=keyboard, parse_mode='HTML')
    await state.set_state(BotStates.confirm)

@dp.callback_query(F.data == "confirm_yes")
async def confirm_yes(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    lang = data.get('lang', 'uz')
    doc_type = data.get('doc_type')
    topic = data.get('topic')
    pages = data.get('pages')
    design = data.get('design')
    user_id = callback.from_user.id
    
    await callback.message.edit_text(get_text(lang, 'generating'))
    
    try:
        content = await generate_content_with_gemini(topic, pages, doc_type, lang)
        
        if doc_type == 'presentation':
            filename = f"presentation_{user_id}.pptx"
            create_presentation(content, design, filename)
        else:
            filename = f"document_{user_id}.docx"
            create_document(content, filename)
        
        file = FSInputFile(filename)
        
        await use_generation(user_id)
        remaining, total = await check_and_reset_limit(user_id)
        
        await callback.message.answer_document(document=file, caption=get_text(lang, 'success').format(remaining=remaining, total=total))
        
        if os.path.exists(filename):
            os.remove(filename)
        
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=get_text(lang, 'back_to_start'), callback_data="back_start")]
        ])
        await callback.message.answer("👍", reply_markup=keyboard)
        
        await state.clear()
        
    except:
        await callback.message.answer(get_text(lang, 'error'))
    
    await callback.answer()

@dp.callback_query(F.data == "confirm_no")
async def confirm_no(callback: types.CallbackQuery, state: FSMContext):
    await state.clear()
    await cmd_start(callback.message, state)
    await callback.answer()

@dp.callback_query(F.data == "back_start")
async def back_to_start(callback: types.CallbackQuery, state: FSMContext):
    await state.clear()
    await cmd_start(callback.message, state)
    await callback.answer()

# ============ MAIN ============
async def main():
    await init_db()
    await dp.start_polling(bot, skip_updates=True)

if __name__ == "__main__":
    asyncio.run(main())
